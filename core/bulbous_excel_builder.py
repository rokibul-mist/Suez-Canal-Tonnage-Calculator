"""
Bulbous Bow Tonnage Calculation Excel Generator
Builds the dedicated 4-Table official calculation workbook for Bulbous Bow Tonnage:
  1. Table 1: Station 20 Transverse Area Calculation
  2. Table 2: Station A Transverse Area Calculation
  3. Table 3: Station B Transverse Area Calculation
  4. Table 4: Bulbous Bow Volume and Tonnage Calculation
Strictly complies with SCA Rules of Navigation and user instructions.
100% pure standard ASCII English encoding.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Styles & Palette
NAVY       = PatternFill("solid", fgColor="1B365D")
HEADER_BLU = PatternFill("solid", fgColor="2B5B84")
SLATE_GRAY = PatternFill("solid", fgColor="455A64")
INPUT_YEL  = PatternFill("solid", fgColor="FFF9C4")
CALC_GRN   = PatternFill("solid", fgColor="E8F5E9")
ALERT_ORG  = PatternFill("solid", fgColor="FFE0B2")
LIGHT_GRAY = PatternFill("solid", fgColor="F5F5F5")
GROUP_HDR  = PatternFill("solid", fgColor="D0E1F9")
ACCENT_GRY = PatternFill("solid", fgColor="E0E0E0")
WHITE      = PatternFill("solid", fgColor="FFFFFF")

TITLE_FONT        = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
SUB_WHITE_BOLD    = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
SECTION_FONT      = Font(name="Calibri", bold=True, size=10, color="1B365D")
HEADER_WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
LABEL_BOLD        = Font(name="Calibri", bold=True, size=9)
BODY_REG          = Font(name="Calibri", size=9)
BODY_BOLD         = Font(name="Calibri", bold=True, size=9)
ITALIC_NOTE       = Font(name="Calibri", italic=True, size=8, color="555555")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
)
THICK_TOP_DOUBLE_BOTTOM = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="medium", color="1B365D"), bottom=Side(style="double", color="1B365D"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

def style_c(ws, r, c, val=None, fill=None, font=None, align=None, border=THIN_BORDER, num_format=None):
    cell = ws.cell(row=r, column=c, value=val)
    if fill: cell.fill = fill
    if font: cell.font = font
    if align: cell.alignment = align
    if border: cell.border = border
    if num_format: cell.number_format = num_format
    return cell

def build_four_tables_bulbous_workbook(output_path):
    """
    Builds the dedicated 4-Table Bulbous Bow Tonnage Workbook:
      - Table 1: Station 20 Transverse Area Calculation
      - Table 2: Station A Transverse Area Calculation
      - Table 3: Station B Transverse Area Calculation
      - Table 4: Bulbous Bow Volume and Tonnage Calculation
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bulbous_Bow_Calculation"
    ws.views.sheetView[0].showGridLines = True

    # 1. Main Title Block
    ws.merge_cells("A1:G1")
    style_c(ws, 1, 1, "SUEZ CANAL TONNAGE - BULBOUS BOW CALCULATION", fill=NAVY, font=TITLE_FONT, align=CENTER)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:G2")
    style_c(ws, 2, 1, "Vessel: M/V TERVEL | Lbp = 186.10 m | B = 32.26 m | D = 17.80 m | Statutory Moorsom Tonnage (V / 2.83)", fill=HEADER_BLU, font=SUB_WHITE_BOLD, align=CENTER)
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:G3")
    style_c(ws, 3, 1, "Procedure: Transverse Simpson 1-4-1 on buttock cuts (C.L., 1000BL, 2000BL) | Longitudinal Simpson 1-4-1 (dx = 1.50 m) + Fore Nose Cap", fill=LIGHT_GRAY, font=ITALIC_NOTE, align=LEFT)
    ws.row_dimensions[3].height = 18

    # -------------------------------------------------------------------------
    # PARAMETER CONFIGURATION BLOCK
    # -------------------------------------------------------------------------
    ws.merge_cells("A5:G5")
    style_c(ws, 5, 1, "1. BULBOUS BOW GEOMETRY & INTEGRATION PARAMETERS", fill=GROUP_HDR, font=SECTION_FONT, align=LEFT)
    ws.row_dimensions[5].height = 20

    cfg_labels = [
        ("A6", "B6", "Total Length of Bulbous Bow (L_bulb):", "C6", 5.000, "0.000", INPUT_YEL, "From FP (Station 20) to foremost nose tip"),
        ("A7", "B7", "Station 20 to Station B Common Interval (dx):", "C7", 1.500, "0.000", INPUT_YEL, "Spacing between St. 20, St. A, and St. B = 1.500 m"),
        ("A8", "B8", "1/3 Longitudinal Multiplier (dx / 3):", "C8", "=C7/3", "0.0000", CALC_GRN, "Simpson 1st Rule longitudinal multiplier = 0.5000 m"),
        ("A9", "B9", "Forward Nose Cap Length (Station B to Tip):", "C9", 2.000, "0.000", INPUT_YEL, "Length of terminal conical / parabolic cap forward of St. B"),
        ("A10", "B10", "Transverse Buttock Spacing (dy):", "C10", 1.000, "0.000", INPUT_YEL, "Cutting spacing: C.L. (0m) to 1000BL (1m) to 2000BL (2m)"),
        ("A11", "B11", "1/3 Transverse Multiplier (dy / 3):", "C11", "=C10/3", "0.0000", CALC_GRN, "Simpson 1st Rule transverse multiplier = 0.3333 m"),
        ("A12", "B12", "Suez Moorsom Conversion Divisor:", "C12", 2.83, "0.00", LIGHT_GRAY, "Statutory divisor: 100 cu. ft. = 2.83 m3 / Suez Ton"),
    ]

    for start_l, end_l, lbl, val_pos, val, fmt, bg, note in cfg_labels:
        r_num = int(start_l[1:])
        ws.merge_cells(f"{start_l}:{end_l}")
        style_c(ws, r_num, 1, lbl, fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
        style_c(ws, r_num, 2, fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
        style_c(ws, r_num, 3, val, fill=bg, font=BODY_BOLD, align=CENTER, num_format=fmt)
        ws.merge_cells(f"D{r_num}:G{r_num}")
        style_c(ws, r_num, 4, note, fill=LIGHT_GRAY, font=ITALIC_NOTE, align=LEFT)
        for c in range(5, 8): style_c(ws, r_num, c, fill=LIGHT_GRAY)
        ws.row_dimensions[r_num].height = 19

    # Helper function to render a single Station Transverse Area Table
    def render_station_table(start_r, title_text, data_rows):
        # Section Title Row
        ws.merge_cells(f"A{start_r}:G{start_r}")
        style_c(ws, start_r, 1, title_text, fill=NAVY, font=HEADER_WHITE_BOLD, align=LEFT)
        for c in range(2, 8): style_c(ws, start_r, c, fill=NAVY)
        ws.row_dimensions[start_r].height = 20

        # Column Headers
        headers = ["Buttock Line", "Half-Breadth Y (m)", "Upper Z (m)", "Lower Z (m)", "Depth D (m)", "Factor", "Product (D x M)"]
        hdr_r = start_r + 1
        for c_idx, h in enumerate(headers, start=1):
            style_c(ws, hdr_r, c_idx, h, fill=SLATE_GRAY, font=SUB_WHITE_BOLD, align=CENTER)
        ws.row_dimensions[hdr_r].height = 19

        data_start = hdr_r + 1
        for idx, (b_name, b_y, z_up, z_low, fac) in enumerate(data_rows):
            curr_r = data_start + idx
            style_c(ws, curr_r, 1, b_name, fill=LIGHT_GRAY, font=BODY_BOLD, align=LEFT)
            style_c(ws, curr_r, 2, b_y, fill=LIGHT_GRAY, font=BODY_REG, align=RIGHT, num_format="0.000")
            style_c(ws, curr_r, 3, z_up, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.000")
            style_c(ws, curr_r, 4, z_low, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.000")
            # Depth D = MAX(0, Upper - Lower)
            style_c(ws, curr_r, 5, f"=MAX(0, C{curr_r}-D{curr_r})", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="0.000")
            style_c(ws, curr_r, 6, fac, fill=LIGHT_GRAY, font=BODY_BOLD, align=CENTER)
            # Product = Depth * Factor
            style_c(ws, curr_r, 7, f"=E{curr_r}*F{curr_r}", fill=CALC_GRN, font=BODY_REG, align=RIGHT, num_format="#,##0.000")
            ws.row_dimensions[curr_r].height = 18

        data_end = data_start + len(data_rows) - 1

        # Sum of Products
        sum_r = data_end + 1
        ws.merge_cells(f"A{sum_r}:F{sum_r}")
        style_c(ws, sum_r, 1, "Sum of Products", fill=ACCENT_GRY, font=LABEL_BOLD, align=LEFT)
        for c in range(2, 7): style_c(ws, sum_r, c, fill=ACCENT_GRY)
        style_c(ws, sum_r, 7, f"=SUM(G{data_start}:G{data_end})", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
        ws.row_dimensions[sum_r].height = 18

        # Half Station Area
        half_r = sum_r + 1
        ws.merge_cells(f"A{half_r}:F{half_r}")
        style_c(ws, half_r, 1, "Half Station Area: (dy / 3) * Sum", fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
        for c in range(2, 7): style_c(ws, half_r, c, fill=LIGHT_GRAY)
        style_c(ws, half_r, 7, f"=$C$11*G{sum_r}", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
        ws.row_dimensions[half_r].height = 18

        # Total Station Area = 2 * Half Area
        tot_r = half_r + 1
        ws.merge_cells(f"A{tot_r}:F{tot_r}")
        style_c(ws, tot_r, 1, "Total Station Area: 2 * Half Area", fill=NAVY, font=HEADER_WHITE_BOLD, align=LEFT)
        for c in range(2, 7): style_c(ws, tot_r, c, fill=NAVY)
        style_c(ws, tot_r, 7, f"=2*G{half_r}", fill=ALERT_ORG, font=Font(name="Calibri", bold=True, size=10, color="1B365D"), align=RIGHT, border=THICK_TOP_DOUBLE_BOTTOM, num_format="#,##0.000")
        ws.row_dimensions[tot_r].height = 20

        return tot_r, f"G{tot_r}"

    # TABLE 1: Station 20 Transverse Area Calculation
    st20_data = [
        ("Centerline (C.L.)", 0.000, 10.700, 0.701, 1),
        ("1000 B.L.",         1.000,  9.223, 1.064, 4),
        ("2000 B.L.",         2.000,  8.388, 1.734, 1),
    ]
    end_r1, cell_area_20 = render_station_table(14, "Station 20 Transverse Area Calculation", st20_data)

    # TABLE 2: Station A Transverse Area Calculation
    stA_data = [
        ("Centerline (C.L.)", 0.000, 8.795, 1.312, 1),
        ("1000 B.L.",         1.000, 8.252, 1.846, 4),
        ("2000 B.L.",         2.000, 7.173, 2.988, 1),
    ]
    end_r2, cell_area_A = render_station_table(end_r1 + 2, "Station A Transverse Area Calculation", stA_data)

    # TABLE 3: Station B Transverse Area Calculation
    stB_data = [
        ("Centerline (C.L.)", 0.000, 8.133, 2.224, 1),
        ("1000 B.L.",         1.000, 7.148, 3.201, 4),
        ("2000 B.L.",         2.000, 0.000, 0.000, 1),
    ]
    end_r3, cell_area_B = render_station_table(end_r2 + 2, "Station B Transverse Area Calculation", stB_data)

    # TABLE 4: Bulbous Bow Volume and Tonnage Calculation
    long_start_r = end_r3 + 2
    ws.merge_cells(f"A{long_start_r}:G{long_start_r}")
    style_c(ws, long_start_r, 1, "Bulbous Bow Volume and Tonnage Calculation", fill=NAVY, font=HEADER_WHITE_BOLD, align=LEFT)
    for c in range(2, 8): style_c(ws, long_start_r, c, fill=NAVY)
    ws.row_dimensions[long_start_r].height = 20

    long_hdrs = ["Station", "Position X (m)", "Transverse Area (m2)", "Factor", "Product (A x M)", "Remarks", ""]
    hdr_r4 = long_start_r + 1
    for c_idx in range(1, 6):
        style_c(ws, hdr_r4, c_idx, long_hdrs[c_idx-1], fill=SLATE_GRAY, font=SUB_WHITE_BOLD, align=CENTER)
    ws.merge_cells(f"F{hdr_r4}:G{hdr_r4}")
    style_c(ws, hdr_r4, 6, "Remarks", fill=SLATE_GRAY, font=SUB_WHITE_BOLD, align=CENTER)
    style_c(ws, hdr_r4, 7, fill=SLATE_GRAY)
    ws.row_dimensions[hdr_r4].height = 19

    long_data = [
        ("Station 20", 0.000, f"={cell_area_20}", 1, "Forward Perpendicular (FP)"),
        ("Station A",  1.500, f"={cell_area_A}",  4, "Intermediate Bulb Station"),
        ("Station B",  3.000, f"={cell_area_B}",  1, "Fore Bulb Station"),
    ]

    r_l_start = hdr_r4 + 1
    for idx, (st_n, px, ar_ref, fac, rem) in enumerate(long_data):
        cr = r_l_start + idx
        style_c(ws, cr, 1, st_n, fill=LIGHT_GRAY, font=BODY_BOLD, align=LEFT)
        style_c(ws, cr, 2, px, fill=LIGHT_GRAY, font=BODY_REG, align=RIGHT, num_format="0.000")
        style_c(ws, cr, 3, ar_ref, fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
        style_c(ws, cr, 4, fac, fill=LIGHT_GRAY, font=BODY_BOLD, align=CENTER)
        style_c(ws, cr, 5, f"=C{cr}*D{cr}", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
        ws.merge_cells(f"F{cr}:G{cr}")
        style_c(ws, cr, 6, rem, fill=LIGHT_GRAY, font=BODY_REG, align=LEFT)
        style_c(ws, cr, 7, fill=LIGHT_GRAY)
        ws.row_dimensions[cr].height = 18

    r_l_end = r_l_start + len(long_data) - 1

    # Sum of Products: Station 20 to Station B
    sum_long_r = r_l_end + 1
    ws.merge_cells(f"A{sum_long_r}:D{sum_long_r}")
    style_c(ws, sum_long_r, 1, "Sum of Products: Station 20 to Station B", fill=ACCENT_GRY, font=LABEL_BOLD, align=LEFT)
    for c in range(2, 5): style_c(ws, sum_long_r, c, fill=ACCENT_GRY)
    ws.merge_cells(f"E{sum_long_r}:G{sum_long_r}")
    style_c(ws, sum_long_r, 5, f"=SUM(E{r_l_start}:E{r_l_end})", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
    for c in range(6, 8): style_c(ws, sum_long_r, c, fill=ACCENT_GRY)
    ws.row_dimensions[sum_long_r].height = 18

    # Main Bulb Body Volume: (dx / 3) * Sum
    body_v_r = sum_long_r + 1
    ws.merge_cells(f"A{body_v_r}:D{body_v_r}")
    style_c(ws, body_v_r, 1, "Main Bulb Body Volume: (dx / 3) * Sum", fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
    for c in range(2, 5): style_c(ws, body_v_r, c, fill=LIGHT_GRAY)
    ws.merge_cells(f"E{body_v_r}:G{body_v_r}")
    style_c(ws, body_v_r, 5, f"=$C$8*E{sum_long_r}", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
    for c in range(6, 8): style_c(ws, body_v_r, c, fill=CALC_GRN)
    ws.row_dimensions[body_v_r].height = 18

    # Forward Nose Cap Volume: Station B to Foremost Tip (1/3 * Area_B * L_nose)
    nose_v_r = body_v_r + 1
    ws.merge_cells(f"A{nose_v_r}:D{nose_v_r}")
    style_c(ws, nose_v_r, 1, "Forward Nose Cap Volume: Station B to Fore Tip (1/3 * Area_B * L_nose)", fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
    for c in range(2, 5): style_c(ws, nose_v_r, c, fill=LIGHT_GRAY)
    ws.merge_cells(f"E{nose_v_r}:G{nose_v_r}")
    style_c(ws, nose_v_r, 5, f"=(1/3)*C{r_l_start+2}*$C$9", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
    for c in range(6, 8): style_c(ws, nose_v_r, c, fill=CALC_GRN)
    ws.row_dimensions[nose_v_r].height = 18

    # Total Bulbous Bow Volume: Main Body + Nose Cap
    tot_v_r = nose_v_r + 1
    ws.merge_cells(f"A{tot_v_r}:D{tot_v_r}")
    style_c(ws, tot_v_r, 1, "Total Bulbous Bow Volume: Main Body + Nose Cap (m3)", fill=GROUP_HDR, font=LABEL_BOLD, align=LEFT)
    for c in range(2, 5): style_c(ws, tot_v_r, c, fill=GROUP_HDR)
    ws.merge_cells(f"E{tot_v_r}:G{tot_v_r}")
    style_c(ws, tot_v_r, 5, f"=E{body_v_r}+E{nose_v_r}", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
    for c in range(6, 8): style_c(ws, tot_v_r, c, fill=CALC_GRN)
    ws.row_dimensions[tot_v_r].height = 19

    # Bulbous Bow Suez Moorsom Tonnage: Total Volume / 2.83
    ton_r = tot_v_r + 1
    ws.merge_cells(f"A{ton_r}:D{ton_r}")
    style_c(ws, ton_r, 1, "Bulbous Bow Suez Moorsom Tonnage: Total Volume / 2.83", fill=NAVY, font=HEADER_WHITE_BOLD, align=LEFT)
    for c in range(2, 5): style_c(ws, ton_r, c, fill=NAVY)
    ws.merge_cells(f"E{ton_r}:G{ton_r}")
    style_c(ws, ton_r, 5, f"=E{tot_v_r}/$C$12", fill=ALERT_ORG, font=Font(name="Calibri", bold=True, size=11, color="1B365D"), align=RIGHT, border=THICK_TOP_DOUBLE_BOTTOM, num_format="#,##0.00")
    for c in range(6, 8): style_c(ws, ton_r, c, fill=ALERT_ORG)
    ws.row_dimensions[ton_r].height = 22

    # Column Widths
    col_widths = {
        "A": 26, # Buttock Line / Station
        "B": 20, # Half-Breadth Y (m) / Position X
        "C": 22, # Upper Z (m) / Area A
        "D": 20, # Lower Z (m) / Factor
        "E": 22, # Depth D (m) / Product
        "F": 28, # Factor / Remarks
        "G": 20, # Product
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    if os.path.dirname(output_path):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
