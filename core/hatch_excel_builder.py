"""
Hatch Coaming Tonnage Calculation Excel Generator
Builds the dedicated calculation workbook for Cargo Hatchways & Coamings.
Includes live Excel formulas and complies with SCA 0.5% (5/1000) Statutory Allowance rule.
100% pure standard ASCII English encoding.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Styles & Palette
NAVY       = PatternFill("solid", fgColor="1B365D")
HEADER_BLU = PatternFill("solid", fgColor="2B5B84")
SLATE_GRAY = PatternFill("solid", fgColor="455A64")
INPUT_YEL  = PatternFill("solid", fgColor="FFF9C4")
CALC_GRN   = PatternFill("solid", fgColor="E8F5E9")
ALERT_ORG  = PatternFill("solid", fgColor="FFE0B2")
LIGHT_GRAY = PatternFill("solid", fgColor="F5F5F5")
GROUP_HDR  = PatternFill("solid", fgColor="D0E1F9")
ACCENT_GRY = PatternFill("solid", fgColor="E0E0E0")
WHITE      = PatternFill("solid", fgColor="FFFFFF")

TITLE_FONT        = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
SUB_WHITE_BOLD    = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
SECTION_FONT      = Font(name="Calibri", bold=True, size=10, color="1B365D")
HEADER_WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
LABEL_BOLD        = Font(name="Calibri", bold=True, size=9)
BODY_REG          = Font(name="Calibri", size=9)
BODY_BOLD         = Font(name="Calibri", bold=True, size=9)
ITALIC_NOTE       = Font(name="Calibri", italic=True, size=8, color="555555")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
)
THICK_TOP_DOUBLE_BOTTOM = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="medium", color="1B365D"), bottom=Side(style="double", color="1B365D"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

def style_c(ws, r, c, val=None, fill=None, font=None, align=None, border=THIN_BORDER, num_format=None):
    cell = ws.cell(row=r, column=c, value=val)
    if fill: cell.fill = fill
    if font: cell.font = font
    if align: cell.alignment = align
    if border: cell.border = border
    if num_format: cell.number_format = num_format
    return cell

def build_hatch_coaming_workbook(output_path, hatches=None, pre_hatch_tonnage=32500.0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hatch_Coamings"
    ws.views.sheetView[0].showGridLines = True

    items = hatches if hatches else [
        {"name": "No. 1 Cargo Hatchway", "length": 15.20, "breadth": 14.00, "height": 1.80},
        {"name": "No. 2 Cargo Hatchway", "length": 17.60, "breadth": 16.00, "height": 1.80},
        {"name": "No. 3 Cargo Hatchway", "length": 17.60, "breadth": 16.00, "height": 1.80},
        {"name": "No. 4 Cargo Hatchway", "length": 17.60, "breadth": 16.00, "height": 1.80},
        {"name": "No. 5 Cargo Hatchway", "length": 17.60, "breadth": 16.00, "height": 1.80},
    ]

    # Title Banner
    ws.merge_cells("A1:G1")
    style_c(ws, 1, 1, "SUEZ CANAL TONNAGE - HATCH COAMING CALCULATION", fill=NAVY, font=TITLE_FONT, align=CENTER)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:G2")
    style_c(ws, 2, 1, "Vessel: M/V TERVEL | SCA Rules of Navigation | Statutory Moorsom Tonnage (V / 2.83)", fill=HEADER_BLU, font=SUB_WHITE_BOLD, align=CENTER)
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:G3")
    style_c(ws, 3, 1, "Volume = Length x Breadth x Height | Moorsom Tonnage = Volume / 2.83 | 0.5% (5/1000) Statutory Allowance Rule", fill=LIGHT_GRAY, font=ITALIC_NOTE, align=LEFT)
    ws.row_dimensions[3].height = 18

    # Parameters block
    ws.merge_cells("A5:G5")
    style_c(ws, 5, 1, "1. HATCH COAMING CONFIGURATION & SUEZ STATUTORY RULES", fill=GROUP_HDR, font=SECTION_FONT, align=LEFT)
    ws.row_dimensions[5].height = 20

    cfg_labels = [
        ("A6", "B6", "Number of Cargo Hatch Covers:", "C6", len(items), "0", INPUT_YEL, "User-configured number of hatchways"),
        ("A7", "B7", "Suez Moorsom Conversion Divisor:", "C7", 2.83, "0.00", LIGHT_GRAY, "Statutory divisor: 100 cu. ft. = 2.83 m3 / Suez Ton"),
        ("A8", "B8", "SCA Statutory Hatch Allowance Rate:", "C8", 0.005, "0.000%", LIGHT_GRAY, "5 per 1000 (0.5%) of Gross Tonnage before hatchway addition"),
        ("A9", "B9", "Pre-Hatch Gross Tonnage (Tons):", "C9", pre_hatch_tonnage, "#,##0.00", INPUT_YEL, "Underdeck + Bulbous Bow + Forecastle + Other Spaces"),
    ]

    for start_l, end_l, lbl, val_pos, val, fmt, bg, note in cfg_labels:
        r_num = int(start_l[1:])
        ws.merge_cells(f"{start_l}:{end_l}")
        style_c(ws, r_num, 1, lbl, fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
        style_c(ws, r_num, 2, fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
        style_c(ws, r_num, 3, val, fill=bg, font=BODY_BOLD, align=CENTER, num_format=fmt)
        ws.merge_cells(f"D{r_num}:G{r_num}")
        style_c(ws, r_num, 4, note, fill=LIGHT_GRAY, font=ITALIC_NOTE, align=LEFT)
        for c in range(5, 8): style_c(ws, r_num, c, fill=LIGHT_GRAY)
        ws.row_dimensions[r_num].height = 19

    # Table Header
    start_r = 11
    ws.merge_cells(f"A{start_r}:G{start_r}")
    style_c(ws, start_r, 1, "Hatch Coaming Tonnage Calculation", fill=NAVY, font=HEADER_WHITE_BOLD, align=LEFT)
    for c in range(2, 8): style_c(ws, start_r, c, fill=NAVY)
    ws.row_dimensions[start_r].height = 20

    headers = ["No.", "Hatchway Description", "Length L (m)", "Breadth B (m)", "Height H (m)", "Volume (m3)", "Tonnage (Suez Tons)"]
    hdr_r = start_r + 1
    for c_idx, h in enumerate(headers, start=1):
        style_c(ws, hdr_r, c_idx, h, fill=SLATE_GRAY, font=SUB_WHITE_BOLD, align=CENTER)
    ws.row_dimensions[hdr_r].height = 19

    data_start = hdr_r + 1
    for idx, h_item in enumerate(items, start=1):
        cr = data_start + idx - 1
        h_name = h_item.get("name") or f"No. {idx} Cargo Hatchway"
        l_val = float(h_item.get("length", 0.0))
        b_val = float(h_item.get("breadth", 0.0))
        h_val = float(h_item.get("height", 0.0))

        style_c(ws, cr, 1, idx, fill=LIGHT_GRAY, font=BODY_BOLD, align=CENTER)
        style_c(ws, cr, 2, h_name, fill=LIGHT_GRAY, font=BODY_BOLD, align=LEFT)
        style_c(ws, cr, 3, l_val, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.000")
        style_c(ws, cr, 4, b_val, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.000")
        style_c(ws, cr, 5, h_val, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.000")
        # Volume = L * B * H
        style_c(ws, cr, 6, f"=C{cr}*D{cr}*E{cr}", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
        # Tonnage = Volume / 2.83
        style_c(ws, cr, 7, f"=F{cr}/$C$7", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.00")
        ws.row_dimensions[cr].height = 18

    data_end = data_start + len(items) - 1

    # Total Aggregate Hatchways
    tot_r = data_end + 1
    ws.merge_cells(f"A{tot_r}:E{tot_r}")
    style_c(ws, tot_r, 1, "Total Aggregate Cargo Hatchways", fill=ACCENT_GRY, font=LABEL_BOLD, align=LEFT)
    for c in range(2, 6): style_c(ws, tot_r, c, fill=ACCENT_GRY)
    style_c(ws, tot_r, 6, f"=SUM(F{data_start}:F{data_end})", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
    style_c(ws, tot_r, 7, f"=SUM(G{data_start}:G{data_end})", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, num_format="#,##0.00")
    ws.row_dimensions[tot_r].height = 19

    # 0.5% Statutory Allowance
    allow_r = tot_r + 1
    ws.merge_cells(f"A{allow_r}:E{allow_r}")
    style_c(ws, allow_r, 1, "Statutory 0.5% Allowance: 0.005 x Pre-Hatch Gross Tonnage", fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
    for c in range(2, 6): style_c(ws, allow_r, c, fill=LIGHT_GRAY)
    style_c(ws, allow_r, 6, f"=G{allow_r}*$C$7", fill=LIGHT_GRAY, font=BODY_REG, align=RIGHT, num_format="#,##0.000")
    style_c(ws, allow_r, 7, f"=$C$9*$C$8", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.00")
    ws.row_dimensions[allow_r].height = 18

    # Excess Hatchway Tonnage Added to Gross Tonnage
    excess_r = allow_r + 1
    ws.merge_cells(f"A{excess_r}:E{excess_r}")
    style_c(ws, excess_r, 1, "Excess Cargo Hatchway Tonnage Added to Gross Tonnage", fill=NAVY, font=HEADER_WHITE_BOLD, align=LEFT)
    for c in range(2, 6): style_c(ws, excess_r, c, fill=NAVY)
    style_c(ws, excess_r, 6, f"=G{excess_r}*$C$7", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="#,##0.000")
    style_c(ws, excess_r, 7, f"=MAX(0, G{tot_r}-G{allow_r})", fill=ALERT_ORG, font=Font(name="Calibri", bold=True, size=11, color="1B365D"), align=RIGHT, border=THICK_TOP_DOUBLE_BOTTOM, num_format="#,##0.00")
    ws.row_dimensions[excess_r].height = 22

    # Column Widths
    col_widths = {
        "A": 8,   # No.
        "B": 28,  # Description
        "C": 18,  # Length (m)
        "D": 18,  # Breadth (m)
        "E": 18,  # Height (m)
        "F": 22,  # Volume (m3)
        "G": 22,  # Tonnage
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    if os.path.dirname(output_path):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
