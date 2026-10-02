"""
Master Excel Workbook Builder for Suez Canal Tonnage Tool
Builds the complete 8-sheet Excel workbook with live formulas, dynamic links, and official SCA layout.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY       = PatternFill("solid", fgColor="1B365D")
HEADER_BLU = PatternFill("solid", fgColor="2B5B84")
SLATE_GRAY = PatternFill("solid", fgColor="455A64")
INPUT_YEL  = PatternFill("solid", fgColor="FFF9C4")
CALC_GRN   = PatternFill("solid", fgColor="E8F5E9")
ALERT_ORG  = PatternFill("solid", fgColor="FFE0B2")
LIGHT_GRAY = PatternFill("solid", fgColor="F5F5F5")
GROUP_HDR  = PatternFill("solid", fgColor="D0E1F9")

TITLE_FONT        = Font(name="Calibri", bold=True, size=13, color="1B365D")
HEADER_WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
SUB_WHITE_BOLD    = Font(name="Calibri", bold=True, color="FFFFFF", size=8)
LABEL_BOLD        = Font(name="Calibri", bold=True, size=9)
BODY_REG          = Font(name="Calibri", size=9)
BODY_BOLD         = Font(name="Calibri", bold=True, size=9)
ITALIC_NOTE       = Font(name="Calibri", italic=True, size=8, color="555555")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
)
DOUBLE_BOTTOM = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="double", color="1B365D"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

def style_c(ws, r, c, val=None, fill=None, font=None, align=None, border=THIN_BORDER, num_format=None):
    cell = ws.cell(row=r, column=c, value=val)
    if fill: cell.fill = fill
    if font: cell.font = font
    if align: cell.alignment = align
    if border: cell.border = border
    if num_format: cell.number_format = num_format
    return cell

def build_complete_master_workbook(output_path, parsed_offsets=None, udt_results=None, bulb_results=None):
    """
    Assembles the official 8-sheet Master Suez Canal Tonnage Workbook.
    """
    # Import existing master generator logic
    import sys, os
    parent_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    if parent_dir not in sys.path:
        sys.path.insert(0, parent_dir)
        
    import build_master_transverse_template as bmtt
    import create_bulbous_bow_module as cbbm
    import build_english_offset_workbook as beow
    import create_full_reference_workbook as cfr

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # 1. Shipyard Lines Table Offsets
    beow.build_lines_table_sheet(wb)

    # 2. Master Transverse Sections Sheet (17 Waterlines)
    area_cell_map = bmtt.build_transverse_sections_17wl_sheet(wb)

    # 3. Master Station Areas Summary & Longitudinal Integration
    bmtt.build_station_areas_integration_sheet(wb, area_cell_map)

    # 4. Official Cover Summary
    cfr.make_cover_sheet(wb)

    # 5. Dedicated Bulbous Bow Sheet
    bulb_tonnage_ref = cbbm.build_bulbous_bow_sheet(wb)

    # 6. Supporting Statutory Sheets
    cfr.make_above_deck_sheet(wb)
    cfr.make_hatchways_sheet(wb)
    cfr.make_deductions_sheet(wb)

    # Remove default blank sheet
    wb.remove(default_sheet)

    # Dynamic Links in Cover Summary
    ws_cov = wb["Cover_Summary"]
    ws_cov["D14"] = "=Station_Areas_Integration!G34"
    ws_cov["D15"] = f"=Bulbous_Bow!{bulb_tonnage_ref}"

    if os.path.dirname(output_path):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
