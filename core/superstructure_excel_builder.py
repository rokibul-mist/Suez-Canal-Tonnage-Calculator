"""
Superstructure & Round Houses Tonnage Calculation Excel Generator
Builds the dedicated calculation workbook for Superstructure and Round Houses enclosed spaces.
Includes live Excel formulas for volume, Moorsom tonnage, additions, and certificate deduction categories.
Complies with SCA Rules of Navigation.
100% pure standard ASCII English encoding.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Styles & Palette
NAVY       = PatternFill("solid", fgColor="1B365D")
HEADER_BLU = PatternFill("solid", fgColor="2B5B84")
SLATE_GRAY = PatternFill("solid", fgColor="455A64")
INPUT_YEL  = PatternFill("solid", fgColor="FFF9C4")
CALC_GRN   = PatternFill("solid", fgColor="E8F5E9")
ALERT_ORG  = PatternFill("solid", fgColor="FFE0B2")
LIGHT_GRAY = PatternFill("solid", fgColor="F5F5F5")
GROUP_HDR  = PatternFill("solid", fgColor="D0E1F9")
ACCENT_GRY = PatternFill("solid", fgColor="E0E0E0")
WHITE      = PatternFill("solid", fgColor="FFFFFF")

TITLE_FONT        = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
SUB_WHITE_BOLD    = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
SECTION_FONT      = Font(name="Calibri", bold=True, size=10, color="1B365D")
HEADER_WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
LABEL_BOLD        = Font(name="Calibri", bold=True, size=9)
BODY_REG          = Font(name="Calibri", size=9)
BODY_BOLD         = Font(name="Calibri", bold=True, size=9)
ITALIC_NOTE       = Font(name="Calibri", italic=True, size=8, color="555555")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
)
THICK_TOP_DOUBLE_BOTTOM = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="medium", color="1B365D"), bottom=Side(style="double", color="1B365D"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

def style_c(ws, r, c, val=None, fill=None, font=None, align=None, border=THIN_BORDER, num_format=None):
    cell = ws.cell(row=r, column=c, value=val)
    if fill: cell.fill = fill
    if font: cell.font = font
    if align: cell.alignment = align
    if border: cell.border = border
    if num_format: cell.number_format = num_format
    return cell

def build_superstructure_workbook(output_path, spaces=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Round_Houses"
    ws.views.sheetView[0].showGridLines = True

    from core.superstructure_engine import DEFAULT_SUPERSTRUCTURE_SPACES
    items = spaces if spaces is not None else DEFAULT_SUPERSTRUCTURE_SPACES

    # Title Banner
    ws.merge_cells("A1:L1")
    style_c(ws, 1, 1, "SUEZ CANAL TONNAGE - ROUND HOUSES & DEDUCTED SPACES", fill=NAVY, font=TITLE_FONT, align=CENTER)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:L2")
    style_c(ws, 2, 1, "SCA Rules of Navigation | Moorsom Tonnage (V / 2.83) | Volume = Qty x L x B x H x Form Factor", fill=HEADER_BLU, font=SUB_WHITE_BOLD, align=CENTER)
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:L3")
    style_c(ws, 3, 1, "Round Houses, Deckhouses, and Enclosed Spaces added to Gross Tonnage; Deductions applied for Net Tonnage", fill=LIGHT_GRAY, font=ITALIC_NOTE, align=LEFT)
    ws.row_dimensions[3].height = 18

    # Table Header at Row 5
    headers = [
        "No.", "Deck Level", "Space Description", "Qty",
        "Length L (m)", "Breadth B (m)", "Height H (m)", "Form Factor k",
        "Volume (m3)", "Action (Gross/Net)", "Deduction Category", "Tonnage (Suez Tons)"
    ]
    ws.row_dimensions[5].height = 24
    for c_idx, h_text in enumerate(headers, start=1):
        style_c(ws, 5, c_idx, h_text, fill=HEADER_BLU, font=HEADER_WHITE_BOLD, align=CENTER)

    start_row = 6
    curr_row = start_row
    for idx, item in enumerate(items, start=1):
        deck = item.get("deck_level", "Upper Deck")
        group = item.get("space_group", f"Space {idx}")
        qty = int(item.get("qty", 1) or 1)
        l = float(item.get("length_m", 0.0) or 0.0)
        b = float(item.get("breadth_m", 0.0) or 0.0)
        h = float(item.get("height_m", 0.0) or 0.0)
        k = float(item.get("form_factor", 1.0) or 1.0)
        action = str(item.get("action", "addition")).strip().lower()
        action_text = "Addition" if action == "addition" else "Deduction"
        cat = item.get("deduction_category", "NONE")

        r = curr_row
        style_c(ws, r, 1, idx, font=BODY_BOLD, align=CENTER)
        style_c(ws, r, 2, deck, font=BODY_REG, align=CENTER)
        style_c(ws, r, 3, group, font=BODY_REG, align=LEFT)
        style_c(ws, r, 4, qty, fill=INPUT_YEL, font=BODY_BOLD, align=CENTER, num_format="0")
        style_c(ws, r, 5, l, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.00")
        style_c(ws, r, 6, b, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.00")
        style_c(ws, r, 7, h, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.00")
        style_c(ws, r, 8, k, fill=INPUT_YEL, font=BODY_REG, align=RIGHT, num_format="0.000")
        # Formula for volume: Qty * L * B * H * k
        style_c(ws, r, 9, f"=D{r}*E{r}*F{r}*G{r}*H{r}", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="0.000")
        style_c(ws, r, 10, action_text, font=BODY_BOLD, align=CENTER)
        style_c(ws, r, 11, cat, font=BODY_REG, align=CENTER)
        # Formula for tonnage: Volume / 2.83
        style_c(ws, r, 12, f"=I{r}/2.83", fill=CALC_GRN, font=BODY_BOLD, align=RIGHT, num_format="0.00")
        ws.row_dimensions[r].height = 20
        curr_row += 1

    end_row = curr_row - 1

    # Totals / Summary Rows
    # Row: Total Round Houses / Superstructures (Gross Tonnage)
    r_add = curr_row
    ws.merge_cells(f"A{r_add}:H{r_add}")
    style_c(ws, r_add, 1, "TOTAL ROUND HOUSES, ETC. (GROSS TONNAGE ADDITION)", fill=GROUP_HDR, font=LABEL_BOLD, align=RIGHT)
    for c in range(2, 9):
        style_c(ws, r_add, c, fill=GROUP_HDR)
    style_c(ws, r_add, 9, f"=SUM(I{start_row}:I{end_row})", fill=GROUP_HDR, font=BODY_BOLD, align=RIGHT, num_format="0.000")
    style_c(ws, r_add, 10, "Addition", fill=GROUP_HDR, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_add, 11, "All Round Houses", fill=GROUP_HDR, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_add, 12, f"=SUM(L{start_row}:L{end_row})", fill=GROUP_HDR, font=BODY_BOLD, align=RIGHT, num_format="0.00")
    ws.row_dimensions[r_add].height = 22
    curr_row += 1

    # Row: Officer's Accommodation Deduction
    r_off = curr_row
    ws.merge_cells(f"A{r_off}:H{r_off}")
    style_c(ws, r_off, 1, "OFFICER'S ACCOMMODATION DEDUCTION", fill=ALERT_ORG, font=LABEL_BOLD, align=RIGHT)
    for c in range(2, 9): style_c(ws, r_off, c, fill=ALERT_ORG)
    style_c(ws, r_off, 9, f'=SUMIF(K{start_row}:K{end_row}, "OFFICER\'S ACCOMMODATION", I{start_row}:I{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.000")
    style_c(ws, r_off, 10, "Deduction", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_off, 11, "Officers", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_off, 12, f'=SUMIF(K{start_row}:K{end_row}, "OFFICER\'S ACCOMMODATION", L{start_row}:L{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.00")
    ws.row_dimensions[r_off].height = 20
    curr_row += 1

    # Row: Crew's Accommodation Deduction
    r_crew = curr_row
    ws.merge_cells(f"A{r_crew}:H{r_crew}")
    style_c(ws, r_crew, 1, "CREW'S ACCOMMODATION DEDUCTION", fill=ALERT_ORG, font=LABEL_BOLD, align=RIGHT)
    for c in range(2, 9): style_c(ws, r_crew, c, fill=ALERT_ORG)
    style_c(ws, r_crew, 9, f'=SUMIF(K{start_row}:K{end_row}, "CREW\'S ACCOMMODATION", I{start_row}:I{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.000")
    style_c(ws, r_crew, 10, "Deduction", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_crew, 11, "Crew", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_crew, 12, f'=SUMIF(K{start_row}:K{end_row}, "CREW\'S ACCOMMODATION", L{start_row}:L{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.00")
    ws.row_dimensions[r_crew].height = 20
    curr_row += 1

    # Row: Used in Working the Ship Deduction
    r_work = curr_row
    ws.merge_cells(f"A{r_work}:H{r_work}")
    style_c(ws, r_work, 1, "USED IN WORKING THE SHIP DEDUCTION", fill=ALERT_ORG, font=LABEL_BOLD, align=RIGHT)
    for c in range(2, 9): style_c(ws, r_work, c, fill=ALERT_ORG)
    style_c(ws, r_work, 9, f'=SUMIF(K{start_row}:K{end_row}, "USED IN WORKING THE SHIP", I{start_row}:I{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.000")
    style_c(ws, r_work, 10, "Deduction", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_work, 11, "Navigation/Safety", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_work, 12, f'=SUMIF(K{start_row}:K{end_row}, "USED IN WORKING THE SHIP", L{start_row}:L{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.00")
    ws.row_dimensions[r_work].height = 20
    curr_row += 1

    # Row: Deduction for Propelling Power
    r_prop = curr_row
    ws.merge_cells(f"A{r_prop}:H{r_prop}")
    style_c(ws, r_prop, 1, "DEDUCTION FOR PROPELLING POWER", fill=ALERT_ORG, font=LABEL_BOLD, align=RIGHT)
    for c in range(2, 9): style_c(ws, r_prop, c, fill=ALERT_ORG)
    style_c(ws, r_prop, 9, f'=SUMIF(K{start_row}:K{end_row}, "DEDUCTION FOR PROPELLING POWER", I{start_row}:I{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.000")
    style_c(ws, r_prop, 10, "Deduction", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_prop, 11, "Propelling/Machinery", fill=ALERT_ORG, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_prop, 12, f'=SUMIF(K{start_row}:K{end_row}, "DEDUCTION FOR PROPELLING POWER", L{start_row}:L{end_row})', fill=ALERT_ORG, font=BODY_BOLD, align=RIGHT, num_format="0.00")
    ws.row_dimensions[r_prop].height = 20
    curr_row += 1

    # Row: Total Deductions
    r_tot_ded = curr_row
    ws.merge_cells(f"A{r_tot_ded}:H{r_tot_ded}")
    style_c(ws, r_tot_ded, 1, "TOTAL DEDUCTIONS", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT)
    for c in range(2, 9): style_c(ws, r_tot_ded, c, fill=ACCENT_GRY)
    style_c(ws, r_tot_ded, 9, f"=I{r_off}+I{r_crew}+I{r_work}+I{r_prop}", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, num_format="0.000")
    style_c(ws, r_tot_ded, 10, "Deduction", fill=ACCENT_GRY, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_tot_ded, 11, "Total", fill=ACCENT_GRY, font=BODY_BOLD, align=CENTER)
    style_c(ws, r_tot_ded, 12, f"=L{r_off}+L{r_crew}+L{r_work}+L{r_prop}", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, num_format="0.00")
    ws.row_dimensions[r_tot_ded].height = 22
    curr_row += 1

    # Row: Net Balance
    r_net = curr_row
    ws.merge_cells(f"A{r_net}:H{r_net}")
    style_c(ws, r_net, 1, "NET TONNAGE BALANCE (ADDITIONS - DEDUCTIONS)", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, border=THICK_TOP_DOUBLE_BOTTOM)
    for c in range(2, 9): style_c(ws, r_net, c, fill=ACCENT_GRY, border=THICK_TOP_DOUBLE_BOTTOM)
    style_c(ws, r_net, 9, f"=I{r_add}-I{r_tot_ded}", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, border=THICK_TOP_DOUBLE_BOTTOM, num_format="0.000")
    style_c(ws, r_net, 10, "Net", fill=ACCENT_GRY, font=BODY_BOLD, align=CENTER, border=THICK_TOP_DOUBLE_BOTTOM)
    style_c(ws, r_net, 11, "Net Balance", fill=ACCENT_GRY, font=BODY_BOLD, align=CENTER, border=THICK_TOP_DOUBLE_BOTTOM)
    style_c(ws, r_net, 12, f"=L{r_add}-L{r_tot_ded}", fill=ACCENT_GRY, font=BODY_BOLD, align=RIGHT, border=THICK_TOP_DOUBLE_BOTTOM, num_format="0.00")
    ws.row_dimensions[r_net].height = 24

    # Column widths
    col_widths = {1: 8, 2: 24, 3: 38, 4: 8, 5: 14, 6: 14, 7: 14, 8: 14, 9: 18, 10: 18, 11: 30, 12: 18}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
