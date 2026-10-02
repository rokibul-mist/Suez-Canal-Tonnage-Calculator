"""
Template Generator for Suez Canal Tonnage Calculator
Generates customized empty Table of Offsets (.xlsx) based on user input parameters.
100% pure standard ASCII English encoding.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY       = PatternFill("solid", fgColor="1B365D")
HEADER_BLU = PatternFill("solid", fgColor="2B5B84")
SLATE_GRAY = PatternFill("solid", fgColor="455A64")
INPUT_YEL  = PatternFill("solid", fgColor="FFF9C4")
LIGHT_GRAY = PatternFill("solid", fgColor="F5F5F5")
GROUP_HDR  = PatternFill("solid", fgColor="D0E1F9")

TITLE_FONT        = Font(name="Calibri", bold=True, size=13, color="1B365D")
HEADER_WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
SUB_WHITE_BOLD    = Font(name="Calibri", bold=True, color="FFFFFF", size=8)
LABEL_BOLD        = Font(name="Calibri", bold=True, size=9)
BODY_REG          = Font(name="Calibri", size=9)
BODY_BOLD         = Font(name="Calibri", bold=True, size=9)
ITALIC_NOTE       = Font(name="Calibri", italic=True, size=8, color="555555")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

DEFAULT_STATIONS = [
    "Transom", "0", "0.5", "1", "1.5", "2", "2.5", "3", "4", "5",
    "6", "7", "8", "9", "10", "11", "12", "13", "14", "15",
    "16", "17", "18", "18.5", "19", "19.5", "20", "A", "B"
]

DEFAULT_WATERLINES = [
    "BL", "500WL", "1000WL", "2000WL", "3000WL", "4000WL", "5000WL", "6000WL",
    "7000WL", "8000WL", "9000WL", "10000WL", "11000WL", "12000WL", "14000WL", "16000WL",
    "Upper Deck Side", "Forecastle Deck Side", "Bulwark Top Line"
]

DEFAULT_BUTTOCKS = [
    "C.L.", "1000 B.L.", "2000 B.L.", "4000 B.L.", "6000 B.L.", "8000 B.L.",
    "10000 B.L.", "12000 B.L.", "14000 B.L.", "Flat of Bottom", "Upper Deck Height",
    "Forecastle Deck Height", "Bulwark Top Height"
]

def style_c(ws, r, c, val=None, fill=None, font=None, align=None, border=THIN_BORDER, num_format=None):
    cell = ws.cell(row=r, column=c, value=val)
    if fill: cell.fill = fill
    if font: cell.font = font
    if align: cell.alignment = align
    if border: cell.border = border
    if num_format: cell.number_format = num_format
    return cell

def generate_offset_template(
    output_path,
    vessel_name="NEW VESSEL",
    lbp=186.10,
    breadth=32.26,
    depth=17.80,
    stations=None,
    waterlines=None,
    buttocks=None
):
    """
    Creates an empty, formatted Table of Offsets (.xlsx) template.
    User can directly fill half-breadths and heights in the yellow cells.
    """
    if not stations:
        stations = DEFAULT_STATIONS
    if not waterlines:
        waterlines = DEFAULT_WATERLINES
    if not buttocks:
        buttocks = DEFAULT_BUTTOCKS

    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Sheet 1: Vessel Particulars
    # -------------------------------------------------------------
    ws_vp = wb.active
    ws_vp.title = "Vessel_Particulars"
    ws_vp.sheet_view.showGridLines = True

    ws_vp.merge_cells("A1:D1")
    style_c(ws_vp, 1, 1, "PRINCIPAL VESSEL PARTICULARS", font=TITLE_FONT, align=LEFT)
    ws_vp.row_dimensions[1].height = 24

    particulars = [
        ("Vessel Name", vessel_name, "Ship identifier"),
        ("Length Between Perpendiculars (Lbp)", lbp, "Meters (m)"),
        ("Breadth Moulded (B)", breadth, "Meters (m)"),
        ("Depth Moulded (D)", depth, "Meters (m)"),
        ("Bulbous Bow Length (L_bulb)", 5.00, "Meters (m) forward of FP"),
        ("Tonnage Divisor", 2.83, "m3 per Suez Moorsom Ton"),
    ]

    for idx, (label, val, note) in enumerate(particulars, start=3):
        style_c(ws_vp, idx, 1, label, fill=LIGHT_GRAY, font=LABEL_BOLD, align=LEFT)
        style_c(ws_vp, idx, 2, val, fill=INPUT_YEL, font=BODY_BOLD, align=RIGHT if isinstance(val, (int, float)) else LEFT)
        style_c(ws_vp, idx, 3, note, fill=LIGHT_GRAY, font=ITALIC_NOTE, align=LEFT)
        ws_vp.row_dimensions[idx].height = 20

    ws_vp.column_dimensions["A"].width = 32
    ws_vp.column_dimensions["B"].width = 22
    ws_vp.column_dimensions["C"].width = 28

    # -------------------------------------------------------------
    # Sheet 2: Lines_Table_Offsets (Empty format ready for input)
    # -------------------------------------------------------------
    ws_lto = wb.create_sheet("Lines_Table_Offsets")
    ws_lto.sheet_view.showGridLines = True

    num_hb = len(waterlines)
    num_ht = len(buttocks)
    total_cols = 1 + num_hb + num_ht + 1  # Station + HB + HT + Station

    # Title
    ws_lto.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    style_c(ws_lto, 1, 1, f"TABLE OF OFFSETS - {vessel_name.upper()} (Unit: mm)", font=TITLE_FONT, align=CENTER)
    ws_lto.row_dimensions[1].height = 24

    ws_lto.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    style_c(ws_lto, 2, 1, "Fill in the half-breadths (Centerline to moulded hull) and heights (Baseline to buttocks/deck) in the yellow cells.", font=ITALIC_NOTE, align=CENTER)
    ws_lto.row_dimensions[2].height = 18

    # Super-Header Row 4
    style_c(ws_lto, 4, 1, "Station", fill=NAVY, font=HEADER_WHITE_BOLD, align=CENTER)
    
    # Half-breadths superheader
    ws_lto.merge_cells(start_row=4, start_column=2, end_row=4, end_column=1 + num_hb)
    style_c(ws_lto, 4, 2, "HALF-BREADTHS (mm) - FROM CENTERLINE TO MOULDED SURFACE", fill=NAVY, font=HEADER_WHITE_BOLD, align=CENTER)
    for c in range(3, 2 + num_hb):
        style_c(ws_lto, 4, c, fill=NAVY)

    # Heights superheader
    ht_start_c = 2 + num_hb
    ht_end_c = 1 + num_hb + num_ht
    ws_lto.merge_cells(start_row=4, start_column=ht_start_c, end_row=4, end_column=ht_end_c)
    style_c(ws_lto, 4, ht_start_c, "HEIGHTS ABOVE BASELINE (mm)", fill=HEADER_BLU, font=HEADER_WHITE_BOLD, align=CENTER)
    for c in range(ht_start_c + 1, ht_end_c + 1):
        style_c(ws_lto, 4, c, fill=HEADER_BLU)

    # Station end superheader
    style_c(ws_lto, 4, total_cols, "Station", fill=NAVY, font=HEADER_WHITE_BOLD, align=CENTER)
    ws_lto.row_dimensions[4].height = 22

    # Row 5: Column Subheaders
    style_c(ws_lto, 5, 1, "Station", fill=SLATE_GRAY, font=SUB_WHITE_BOLD, align=CENTER)
    for i, wl in enumerate(waterlines, start=2):
        style_c(ws_lto, 5, i, wl, fill=HEADER_BLU, font=SUB_WHITE_BOLD, align=CENTER)
    for j, bt in enumerate(buttocks, start=ht_start_c):
        style_c(ws_lto, 5, j, bt, fill=NAVY, font=SUB_WHITE_BOLD, align=CENTER)
    style_c(ws_lto, 5, total_cols, "Station", fill=SLATE_GRAY, font=SUB_WHITE_BOLD, align=CENTER)
    ws_lto.row_dimensions[5].height = 26

    # Data Rows
    start_r = 6
    for idx, st_name in enumerate(stations):
        curr_r = start_r + idx
        # Station column
        style_c(ws_lto, curr_r, 1, st_name, fill=LIGHT_GRAY, font=LABEL_BOLD, align=CENTER)
        # Editable Half-breadths
        for c in range(2, 2 + num_hb):
            style_c(ws_lto, curr_r, c, "-", fill=INPUT_YEL, font=BODY_REG, align=CENTER)
        # Editable Heights
        for c in range(ht_start_c, ht_end_c + 1):
            style_c(ws_lto, curr_r, c, "-", fill=INPUT_YEL, font=BODY_REG, align=CENTER)
        # End Station column
        style_c(ws_lto, curr_r, total_cols, st_name, fill=LIGHT_GRAY, font=LABEL_BOLD, align=CENTER)
        ws_lto.row_dimensions[curr_r].height = 19

    # Widths
    ws_lto.column_dimensions["A"].width = 11
    ws_lto.column_dimensions[get_column_letter(total_cols)].width = 11
    for c in range(2, total_cols):
        ws_lto.column_dimensions[get_column_letter(c)].width = 12

    if os.path.dirname(output_path):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
